from io import BytesIO
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange

from .models import Column, Row, SourcePosition, Table, XlsxOptions, XlsxSheetOptions
from .xlsx_values import adapt_cell


def validate_xlsx_options(options: XlsxOptions) -> None:
    """只在公开入口调用；行号基于原工作表，区域只接受封闭 A1 矩形。"""
    if not isinstance(options, XlsxOptions):
        raise ValueError("xlsx must be XlsxOptions or None")
    if options.sheets is None:
        return
    if not isinstance(options.sheets, tuple) or not options.sheets:
        raise ValueError("xlsx.sheets must be a nonempty tuple or None")
    names = set()
    for config in options.sheets:
        if not isinstance(config, XlsxSheetOptions):
            raise ValueError("xlsx.sheets items must be XlsxSheetOptions")
        if not isinstance(config.sheet, str) or not config.sheet.strip() or config.sheet in names:
            raise ValueError("sheet names must be nonempty and unique")
        names.add(config.sheet)
        top, bottom = 1, 1048576
        if config.cell_range is not None:
            if not isinstance(config.cell_range, str) or not re.fullmatch(
                r"\$?[A-Za-z]{1,3}\$?[1-9][0-9]*(?::\$?[A-Za-z]{1,3}\$?[1-9][0-9]*)?",
                config.cell_range,
            ):
                raise ValueError("cell_range must be a closed A1 rectangle")
            left, top, right, bottom = CellRange(config.cell_range).bounds
            if not (1 <= left <= right <= 16384 and 1 <= top <= bottom <= 1048576):
                raise ValueError("cell_range is reversed or outside XLSX bounds")
        if config.header_rows is not None:
            header = config.header_rows
            if not isinstance(header, tuple) or not header or any(
                type(row) is not int or not top <= row <= bottom for row in header
            ):
                raise ValueError("header_rows must be a nonempty tuple of rows within the range")
            if any(current != previous + 1 for previous, current in zip(header, header[1:])):
                raise ValueError("header_rows must be consecutive and increasing")


def read_xlsx(data, source, context, options, warnings: list[str]) -> list[Table]:
    """两份工作簿视图对应同一输入字节；缓存视图只交给值适配器，不进行计算。"""
    book = load_workbook(BytesIO(data), data_only=False)
    try:
        cached = load_workbook(BytesIO(data), data_only=True)
        try:
            selected = None if options is None or options.sheets is None else {
                item.sheet: item for item in options.sheets
            }
            if selected is not None:
                missing = set(selected) - set(book.sheetnames)
                if missing:
                    raise ValueError(f"XLSX sheets not found: {sorted(missing)}")
            tables = []
            for sheet in book.worksheets:
                if selected is not None:
                    if sheet.title not in selected:
                        continue
                    config = selected[sheet.title]
                else:
                    if sheet.sheet_state != "visible":
                        continue
                    config = XlsxSheetOptions(sheet.title)
                populated = [cell for row in sheet.iter_rows() for cell in row
                             if cell.value is not None and str(cell.value).strip()]
                if not populated and selected is None:
                    continue
                if not populated:
                    raise ValueError(f"XLSX sheet {sheet.title}: empty selected sheet")
                bounds = CellRange(config.cell_range).bounds if config.cell_range else (
                    min(cell.column for cell in populated), min(cell.row for cell in populated),
                    max(cell.column for cell in populated), max(cell.row for cell in populated),
                )
                if config.cell_range is None:
                    for merged in sheet.merged_cells.ranges:
                        anchor = sheet.cell(merged.min_row, merged.min_col)
                        if anchor.value is not None and str(anchor.value).strip():
                            bounds = (min(bounds[0], merged.min_col), min(bounds[1], merged.min_row),
                                      max(bounds[2], merged.max_col), max(bounds[3], merged.max_row))
                tables.append(_read_sheet(sheet, cached[sheet.title], source,
                                          config.context or context, bounds, config.header_rows, warnings))
            if not tables:
                raise ValueError("XLSX has no nonempty visible sheets")
            return tables
        finally:
            cached.close()
    finally:
        book.close()


def _read_sheet(sheet, cached_sheet, source, context, bounds, header_rows, warnings):
    left, top, right, bottom = bounds
    location = f"XLSX sheet {sheet.title}"
    columns, header_rows, merges = _read_header(
        sheet, cached_sheet, source, bounds, header_rows, warnings)
    rows = []
    for row in sheet.iter_rows(min_row=header_rows[-1] + 1, max_row=bottom, min_col=left, max_col=right):
        values = [_read_cell(sheet, cached_sheet, source, cell.row, cell.column, merges, warnings)
                  for cell in row]
        if not any(cell.text.strip() for cell in values):
            continue
        rows.append(Row(values, columns, context,
                        _position(source, sheet.title, row[0].row, row[0].row, left, right)))
    if not rows:
        raise ValueError(f"{location}: no data records after header row {header_rows[-1]}")
    return Table(columns, rows, context,
                 _position(source, sheet.title, header_rows[0], rows[-1].source.row_end, left, right))


def _read_header(sheet, cached_sheet, source, bounds, header_rows, warnings):
    """按真实合并关系展开完整列名，拒绝空层、重名及表头/数据跨界。"""
    left, top, right, bottom = bounds
    location = f"XLSX sheet {sheet.title}"
    merges = _selected_merges(sheet, (left, header_rows[0] if header_rows else top, right, bottom))
    if header_rows is None:
        header_rows = next(((row[0].row,) for row in sheet.iter_rows(
            min_row=top, max_row=bottom, min_col=left, max_col=right)
            if any(cell.value is not None and str(cell.value).strip() for cell in row)), None)
    if header_rows is None:
        raise ValueError(f"{location}: header not found in selected range")
    if header_rows[-1] > bottom:
        raise ValueError(f"{location}: header rows outside selected range")
    for merged in set(merges.values()):
        if merged.min_row <= header_rows[-1] < merged.max_row:
            raise ValueError(f"{location}: merged range {merged} crosses header/data boundary")
    levels = []
    for row in header_rows:
        values = [_read_cell(sheet, cached_sheet, source, row, col, merges, warnings, header=True)
                  for col in range(left, right + 1)]
        if not any(cell.text.strip() for cell in values):
            raise ValueError(f"{location}: header row {row} is empty")
        levels.append(values)
    columns = []
    for cells in zip(*levels):
        labels = []
        seen_merges = set()
        for cell in cells:
            merged = cell.source.merged_range
            if cell.text.strip() and (merged is None or merged not in seen_merges):
                labels.append(cell.text)
            if merged is not None:
                seen_merges.add(merged)
        name = " / ".join(labels)
        if not name.strip():
            raise ValueError(f"{location}: empty column name at {cells[0].source.coordinate}")
        columns.append(Column(name, cells[0].source, [cell.source for cell in cells]))
    if len({column.name for column in columns}) != len(columns):
        raise ValueError(f"{location}: duplicate column names at header rows {header_rows}")
    return columns, header_rows, merges


def _selected_merges(sheet, bounds):
    """保留选中结构的真实覆盖关系；切断锚点或覆盖范围时不能猜测其标签。"""
    left, top, right, bottom = bounds
    merges = {}
    for merged in sheet.merged_cells.ranges:
        if merged.min_col > right or merged.max_col < left or merged.min_row > bottom or merged.max_row < top:
            continue
        if not (left <= merged.min_col <= merged.max_col <= right
                and top <= merged.min_row <= merged.max_row <= bottom):
            raise ValueError(f"XLSX sheet {sheet.title}: selected range cuts merged range {merged}")
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                merges[row, col] = merged
    return merges


def _read_cell(sheet, cached_sheet, source, row, col, merges, warnings, *, header=False):
    """覆盖格原值始终来自自身；数据区仅传播文本共同标签，并明示其合并归属。"""
    position = _position(source, sheet.title, row, row, col, col)
    merged = merges.get((row, col))
    if merged is not None:
        position.merged_anchor = f"{get_column_letter(merged.min_col)}{merged.min_row}"
        position.merged_range = str(merged)
    result = adapt_cell(sheet.cell(row, col), cached_sheet.cell(row, col), position, warnings)
    if merged is None:
        return result
    anchor = sheet.cell(merged.min_row, merged.min_col)
    covered = position.coordinate != position.merged_anchor
    if header:
        if covered:
            result.text = adapt_cell(anchor, cached_sheet.cell(anchor.row, anchor.column), position, warnings).text
    elif anchor.data_type == "s":
        if covered and anchor.value is not None and anchor.value.strip():
            result.text = f"{anchor.value}（合并共同标签，锚点 {position.merged_anchor}，范围 {merged}）"
    elif anchor.value is not None:
        if covered:
            result.text = f"（合并覆盖，值见锚点 {position.merged_anchor}，范围 {merged}；非本行独立值）"
        else:
            result.text += f"（合并值，仅锚点记录，范围 {merged}）"
    return result


def _position(source, sheet, top, bottom, left, right):
    start, end = f"{get_column_letter(left)}{top}", f"{get_column_letter(right)}{bottom}"
    return SourcePosition(source, top, bottom, left, right, sheet=sheet,
                          coordinate=start if start == end else f"{start}:{end}")
