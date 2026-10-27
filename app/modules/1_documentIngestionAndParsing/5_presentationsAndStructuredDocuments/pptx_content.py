from dataclasses import replace
from html import escape
from io import BytesIO
import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .models import ContentSource, ConversionResult, ReviewIssue


def extract_pptx(data, source):
    """按原页和对象图层顺序提取；对象位置与图层顺序都不是已核验的阅读顺序。"""
    presentation = Presentation(BytesIO(data))
    if not presentation.slides:
        raise ValueError("Presentation contains no slides")
    writer = _Writer()
    for number, slide in enumerate(presentation.slides, 1):
        title_shape = slide.shapes.title
        page = ContentSource(
            kind="slide", slide_number=number,
            slide_title=title_shape.text if title_shape is not None else None,
        )
        writer.add(f"# 第 {number} 页\n\n", page)
        writer.review(
            "reading_order_unverified",
            "Objects retain file z-order; reading order, columns and condition ownership require original-slide review.",
            page,
        )
        for shape in slide.shapes:
            _extract_shape(writer, shape, page, (), title_shape.shape_id if title_shape is not None else None)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                writer.add("## 演讲者备注\n\n", replace(page, kind="notes"))
                _write_paragraphs(writer, notes, replace(page, kind="notes_paragraph"))
    return ConversionResult(
        "".join(writer.parts), source, sources=writer.sources,
        review_issues=writer.review_issues,
    )


def _extract_shape(writer, shape, page, parent_path, title_id):
    position = replace(
        page, kind="object", markdown_start=None, markdown_end=None,
        shape_path=parent_path + (shape.shape_id,), shape_name=shape.name,
        bounds_emu=(shape.left, shape.top, shape.width, shape.height),
    )
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        position.kind = "group"
        writer.sources.append(position)
        writer.review("group_relationships_unverified", "Group meaning and child coordinates require original-slide review.", position)
        for child in shape.shapes:
            _extract_shape(writer, child, page, position.shape_path, None)
    elif shape.has_table:
        position.kind = "table"
        _write_table(writer, shape.table, position)
    elif shape.has_text_frame:
        position.kind = "title" if not parent_path and shape.shape_id == title_id else "text"
        if position.kind == "title":
            writer.sources.append(position)
        else:
            writer.add(f"### 文本对象 {shape.shape_id}\n\n", position)
        _write_paragraphs(writer, shape.text_frame, position, title=position.kind == "title")
        if shape.shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM):
            writer.review("shape_relationships_unverified", "Shape text was extracted; shape meaning, arrows and relationships were not verified.", position)
    else:
        position.kind = {
            MSO_SHAPE_TYPE.PICTURE: "picture",
            MSO_SHAPE_TYPE.LINKED_PICTURE: "picture",
            MSO_SHAPE_TYPE.LINE: "connector",
            MSO_SHAPE_TYPE.CHART: "chart",
        }.get(shape.shape_type, "unsupported_object")
        writer.sources.append(position)
        writer.review(
            f"{position.kind}_unverified",
            "Visual content, labels, arrow direction and relationships were not extracted; review the original object.",
            position,
        )


def _write_paragraphs(writer, frame, position, *, title=False):
    for number, paragraph in enumerate(frame.paragraphs, 1):
        if not paragraph.text.strip():
            continue
        source = replace(
            position, kind="title_paragraph" if title else position.kind,
            markdown_start=None, markdown_end=None, paragraph_number=number, level=paragraph.level,
        )
        text = _markdown_text(paragraph.text)
        body = f"## {text}\n\n" if title else f"{'  ' * paragraph.level}- {text}\n\n"
        writer.add(body, source)


def _write_table(writer, table, position):
    """HTML 行列保留普通及合并单元格，不把第一行臆定为表头。"""
    writer.add(f"### 表格对象 {position.shape_path[-1]}\n\n", position)
    writer.add("<table>\n")
    merged = False
    for row_number, row in enumerate(table.rows, 1):
        writer.add("<tr>")
        for column_number, cell in enumerate(row.cells, 1):
            if cell.is_spanned:
                continue
            attributes = ""
            if cell.is_merge_origin:
                merged = True
                attributes = f' rowspan="{cell.span_height}" colspan="{cell.span_width}"'
            text = escape(cell.text).replace("\v", "<br>").replace("\n", "<br>")
            writer.add(
                f"<td{attributes}>{text}</td>",
                replace(
                    position, kind="table_cell", markdown_start=None, markdown_end=None,
                    row_number=row_number, column_number=column_number,
                ),
            )
        writer.add("</tr>\n")
    writer.add("</table>\n\n")
    if merged:
        writer.review("merged_table_unverified", "Cell spans are retained; header and condition ownership require review.", position)


def _markdown_text(text):
    text = re.sub(r"([\\`*_{}\[\]<>#+.!|~-])", r"\\\1", text)
    return text.replace("\v", "<br>").replace("\n", "<br>")


class _Writer:
    """正文及字符区间同时生成，非文本对象只有来源，不伪造正文区间。"""

    def __init__(self):
        self.parts = []
        self.length = 0
        self.sources = []
        self.review_issues = []

    def add(self, text, source=None):
        if source is not None:
            source.markdown_start = self.length
            source.markdown_end = self.length + len(text)
            self.sources.append(source)
        self.parts.append(text)
        self.length += len(text)

    def review(self, code, message, source):
        self.review_issues.append(ReviewIssue(code, message, source))
