"""ZIP #41 固定输入及模拟服务产物；模拟产物不代表真实识别质量。"""

import hashlib
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile, ZipInfo


SAMPLES = Path(__file__).parent
SERVICE_MARKDOWN = "# Service original\n\nKeep **formatting**, order and trailing spaces.  \n"
BLOCKS = [{"type": "text", "text": "Different structured text", "page_idx": 0,
           "bbox": [10, 20, 300, 400]}]


def archive(name, entries, compression=ZIP_DEFLATED):
    with ZipFile(SAMPLES / name, "w") as output:
        for path, data in entries:
            info = ZipInfo(path, date_time=(2026, 10, 28, 0, 0, 0))
            info.compress_type = compression
            output.writestr(info, data)


def main():
    pdf = [("docs/manual.pdf", (SAMPLES / "pdf_digital_teaching.pdf").read_bytes()),
           ("after.txt", "Sibling survives")]
    office = [("docs/report.docx", (SAMPLES / "docx_sections.docx").read_bytes()),
              ("slides/deck.pptx", (SAMPLES / "pptx_refund.pptx").read_bytes()),
              ("tables/costs.xlsx", (SAMPLES / "xlsx_travel.xlsx").read_bytes())]
    web = [("web/page.html", b"<html><body><h1>Remote article</h1></body></html>"),
           ("web/alias.htm", b"<html><body>Alias HTML</body></html>"),
           ("web/archive.mhtml", (SAMPLES / "mhtml_start_shorter.mhtml").read_bytes()),
           ("web/fallback.mht", (SAMPLES / "mhtml_character_tie.mhtml").read_bytes())]
    archive("zip_document_pdf.zip", pdf)
    archive("zip_document_office.zip", office)
    archive("zip_document_web.zip", web)
    archive("zip_documents.zip", pdf + office + web)
    result = [("output/full.md", SERVICE_MARKDOWN),
              ("output/generated_content_list.json", json.dumps(BLOCKS)),
              ("output/generated_content_list_v2.json", '[[{"type":"text","content":"V2"}]]'),
              ("output/layout.json", '{"pdf_info":[],"_backend":"fixture"}')]
    archive("zip_document_result.zip", result)
    archive("zip_document_html_result.zip", [("full.md", "# Remote article\n"),
                                             ("main.html", "<h1>Remote article</h1>")])
    archive("zip_document_result_no_md.zip", result[1:])
    archive("zip_document_result_bad_json.zip", [("full.md", "Partial body"),
                                               ("bad_content_list.json", "{")])
    archive("zip_document_result_v2_only.zip", [result[0], result[2]])
    archive("zip_document_result_bad_list.zip", [("full.md", "Partial body"),
                                               ("bad_content_list.json", '[{"page_idx":0}]')])
    archive("zip_document_result_many.zip", result + [("a", "1"), ("b", "2")])
    archive("zip_document_result_expanded.zip", result + [("image.png", b"x" * 200_000)])
    archive("zip_document_result_download.zip", result + [("image.png", b"x" * 200_000)], ZIP_STORED)
    archive("zip_document_result_blank.zip", [("full.md", " \n"), result[1]])
    archive("zip_document_result_missing_page.zip", [("full.md", SERVICE_MARKDOWN),
            ("generated_content_list.json", '[{"type":"text","text":"Body"}]')])
    build_deadline_sample()
    build_live_samples()


def build_deadline_sample():
    """前一文档可成功，后一文档截止后仍有文档、目录和确定的系统元数据。"""
    pdf = (SAMPLES / "pdf_digital_teaching.pdf").read_bytes()
    archive("zip_document_deadline.zip", [
        ("before.txt", "Completed text"), ("completed.pdf", pdf), ("stopped.pdf", pdf),
        ("after.txt", "Must remain unprocessed"), ("after.pdf", pdf), ("folder/", b""),
        ("__MACOSX/file", b"metadata"), (".DS_Store", b""), ("._manual.pdf", b""),
    ])


def build_live_samples():
    """验收事实来自构造输入，不从服务响应反推期望。"""
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Maintenance Start"
    slide.placeholders[1].text = "Disconnect power before maintenance."
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Maintenance Schedule"
    table = slide.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
    for row, values in zip(table.rows, [("Equipment", "Interval"), ("Pump", "Monthly"), ("Valve", "Quarterly")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    deck.save(SAMPLES / "zip_document_live_slides.pptx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Travel"
    sheet.append(["City", "Allowance", None])
    sheet.merge_cells("B1:C1")
    sheet.append([None, "Hotel", "Meals"])
    sheet.append(["Beijing", 500, 100])
    sheet.append(["Shanghai", 450, 100])
    sheet.append(["Total", "=SUM(B3:B4)"])
    second = workbook.create_sheet("Equipment")
    second.append(["Equipment", "Interval"])
    second.append(["Pump", "Monthly"])
    workbook.save(SAMPLES / "zip_document_live_workbook.xlsx")

    html = ("<!doctype html><html><head><meta charset='utf-8'><title>Maintenance Rules</title></head>"
            "<body><article><h1>Maintenance Rules</h1><p>Disconnect power before maintenance.</p>"
            "<table><tr><th>Equipment</th><th>Interval</th></tr><tr><td>Pump</td><td>Monthly</td></tr>"
            "<tr><td>Valve</td><td>Quarterly</td></tr></table><p>Restart after inspection.</p></article></body></html>")
    (SAMPLES / "zip_document_live_page.html").write_text(html, encoding="utf-8", newline="\n")
    mhtml = ("MIME-Version: 1.0\nContent-Type: multipart/related; boundary=live; start=\"<main>\"\n\n"
             "--live\nContent-Type: text/html; charset=utf-8\nContent-ID: <main>\n\n" + html
             + "\n--live--\n")
    (SAMPLES / "zip_document_live_page.mhtml").write_bytes(mhtml.encode())

    cases = {
        "pdf_digital": ("pdf_digital_teaching.pdf", ["Equipment Maintenance", "Disconnect power before maintenance.", "Inspect every month."], [], [1]),
        "pdf_scan": ("pdf_scanned_maintenance_table.pdf", ["Disconnect power before maintenance.", "Restart after inspection.", "Temperature must not exceed 60 C."],
                     [["Equipment", "Interval"], ["Pump", "Monthly"], ["Valve", "Quarterly"]], [1, 2]),
        "pdf_mixed": ("pdf_mixed_maintenance.pdf", ["Digital Intake", "Asset code: PUMP-A17.", "Disconnect power before maintenance.", "Digital Release", "Torque setting: 12 Nm.", "Temperature must not exceed 60 C."],
                      [["Equipment", "Interval"], ["Pump", "Monthly"], ["Valve", "Quarterly"]], [1, 2, 3, 4]),
        "docx": ("docx_sections.docx", ["设备维护", "维护前必须断电。", "检查后重新启动。"], [["设备", "周期"], ["泵", "每月"], ["阀门", "每季度"]], []),
        "pptx": ("zip_document_live_slides.pptx", ["Maintenance Start", "Disconnect power before maintenance.", "Maintenance Schedule"],
                 [["Equipment", "Interval"], ["Pump", "Monthly"], ["Valve", "Quarterly"]], []),
        "xlsx": ("zip_document_live_workbook.xlsx", ["Beijing", "Shanghai", "Equipment", "Pump"],
                 [["Beijing", "500", "100"], ["Shanghai", "450", "100"], ["Pump", "Monthly"]], []),
        "html": ("zip_document_live_page.html", ["Maintenance Rules", "Disconnect power before maintenance.", "Restart after inspection."],
                 [["Equipment", "Interval"], ["Pump", "Monthly"], ["Valve", "Quarterly"]], []),
        "mhtml": ("zip_document_live_page.mhtml", ["Maintenance Rules", "Disconnect power before maintenance.", "Restart after inspection."],
                  [["Equipment", "Interval"], ["Pump", "Monthly"], ["Valve", "Quarterly"]], []),
    }
    manifest = {}
    for case, (name, facts, rows, pages) in cases.items():
        payload = (SAMPLES / name).read_bytes()
        zip_name = f"zip_document_live_{case}.zip"
        archive(zip_name, [(name, payload)])
        manifest[case] = {"archive": zip_name, "member": name, "source_sha256": hashlib.sha256(payload).hexdigest(),
                          "archive_sha256": hashlib.sha256((SAMPLES / zip_name).read_bytes()).hexdigest(),
                          "facts": facts, "table_rows": rows, "pdf_pages": pages}
    (SAMPLES / "zip_document_live_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")


if __name__ == "__main__":
    main()
