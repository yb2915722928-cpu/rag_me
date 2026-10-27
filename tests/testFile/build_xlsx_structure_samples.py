"""构造 #17 固定结构样本；仅使用 openpyxl，不导入生产解析器。"""
from pathlib import Path
from openpyxl import Workbook

ROOT = Path(__file__).parent
book = Workbook()
s = book.active
s.title = "层级"
s.append(["说明：显式从第2行读取"])
s.append(["对象", "标准", None, "备注"])
s.append(["城市", "住宿（元／晚）", "餐补（元／天）", None])
s.append(["北京", 500, 100, "教学"])
s.append(["上海", 450, 100, None])
book.save(ROOT / "xlsx_structure_layers.xlsx")

book = Workbook()
s = book.active
s.title = "合并"
s.append(["地区", "差旅标准", None, "备注"])
s.append([None, "住宿", "餐补", None])
s.append([None, "元／晚", "元／天", None])
s.append(["华北", 500, 100, "共同说明"])
s.append([None, 450, 100, None])
s.append([None, 400, 80, None])
s.append(["华东", "横向共同标签", None, None])
s.merge_cells("A1:A3")
s.merge_cells("B1:C1")
s.merge_cells("D1:D3")
s.merge_cells("A4:A5")
s.merge_cells("D4:D5")
s.merge_cells("B7:D7")
book.save(ROOT / "xlsx_structure_merges.xlsx")

book = Workbook()
s = book.active
s.title = "测量"
s.append(["对象", "金额（元）", "公式", "标志"])
s.append(["甲", 900, "=B2+1", True])
s.append(["乙", None, None, None])
s.append(["丙", None, None, None])
s.merge_cells("B2:B3")
s.merge_cells("C2:C3")
s.merge_cells("D2:D3")
book.save(ROOT / "xlsx_structure_measurements.xlsx")

for suffix, records, merges in (
    ("duplicate", [["标准", "标准"], ["金额", "金额"], [1, 2]], []),
    ("blank", [["标准", None], ["金额", None], [1, 2]], []),
    ("cross_boundary", [["对象", "金额"], ["北京", 1], [None, 2]], ["A1:A2"]),
    ("empty_level", [["对象", "金额"], [None, None], ["北京", 1]], []),
    ("vertical_data", [["标签"], ["共同"], [None], [None]], ["A2:A4"]),
):
    book = Workbook()
    book.active.title = "结构"
    for record in records:
        book.active.append(record)
    for area in merges:
        book.active.merge_cells(area)
    book.save(ROOT / f"xlsx_structure_{suffix}.xlsx")


book = Workbook()
book.active.title = "有效"
book.active.append(["对象", "金额"])
book.active.append(["北京", 500])
s = book.create_sheet("跨界")
s.append(["对象", "金额"])
s.append([None, 100])
s.merge_cells("A1:A2")
book.save(ROOT / "xlsx_structure_atomic.xlsx")
