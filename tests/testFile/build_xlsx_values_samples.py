"""构造固定输入；OOXML v 为人为预设缓存，不是办公软件重算结果。"""

from datetime import datetime, time
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook


def build():
    book = Workbook()
    sheet = book.active
    sheet.title = "值"
    sheet.append(["项目", "值"])
    samples = [
        ("缓存零", "=1-1", "General"),
        ("无缓存", "=1+1", "General"),
        ("缓存数字", "=40+2", "General"),
        ("缓存错误", "=1/0", "0.00"),
        ("空白", None, "General"),
        ("零", 0, "General"),
        ("错误", "#VALUE!", "0.00"),
        ("整数", 1234, "#,##0"),
        ("小数", -1234.5, "#,##0.00"),
        ("小数补零", 12.5, "0.00"),
        ("百分比", 0.2, "0%"),
        ("百分比小数", 0.125, "0.00%"),
        ("金额", 1234.5, '"¥"#,##0.00'),
        ("美元", -12.5, "$0.00"),
        ("日期", datetime(2026, 10, 27), "yyyy-mm-dd"),
        ("日期时间", datetime(2026, 10, 27, 13, 4, 5), "yyyy-mm-dd h:mm:ss"),
        ("时间", time(13, 4, 5), "h:mm:ss"),
        ("缓存百分比", "=1/5", "0%"),
        ("缓存日期", "=DATE(2026,10,27)", "yyyy-mm-dd"),
        ("自定义", 12.5, '0.0" widgets"'),
        ("布尔", False, "0.00"),
        ("文本", "00123", "0.00"),
        ("缓存自定义", "=25/2", '0.0" widgets"'),
        ("缓存金额", "=2500/2", '"¥"#,##0.00'),
        ("缓存文本", '=IF(1,"00123","")', "0.00"),
        ("缓存布尔", "=1=2", "0.00"),
        ("自定义日期", datetime(2026, 10, 27), 'yyyy"年"m"月"d"日"'),
        ("自定义倍率", 12000, "#,##0,"),
        ("小数舍入", 2.675, "0.00"),
        ("整数舍入", 2.5, "0"),
        ("负数舍入", -2.5, "0"),
    ]
    for label, value, number_format in samples:
        sheet.append([label, value])
        sheet.cell(sheet.max_row, 2).number_format = number_format
    buffer = BytesIO()
    book.save(buffer)
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    caches = {
        "B2": ("n", "0"), "B4": ("n", "42"), "B5": ("e", "#DIV/0!"),
        "B19": ("n", "0.2"), "B20": ("n", "46322"),
        "B24": ("n", "12.5"), "B25": ("n", "1250"),
        "B26": ("str", "00123"), "B27": ("b", "0"),
    }
    target = Path(__file__).with_name("xlsx_values.xlsx")
    with ZipFile(buffer) as original, ZipFile(target, "w") as output:
        for info in original.infolist():
            data = original.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(data)
                for cell in root.iter(f"{{{namespace}}}c"):
                    if cell.attrib["r"] in caches:
                        kind, value = caches[cell.attrib["r"]]
                        cell.set("t", kind)
                        cell.find(f"{{{namespace}}}v").text = value
                data = ET.tostring(root, encoding="utf-8")
            output.writestr(info, data)


if __name__ == "__main__":
    build()
