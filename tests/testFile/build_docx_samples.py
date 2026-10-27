"""使用 python-docx 1.2.0 构造补充样本；历史教学副本不在此重建。"""

from copy import copy
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SAMPLES = Path(__file__).parent


def main():
    document = Document()
    document.add_paragraph("文档说明。")
    document.add_heading("设备维护", level=1)
    document.add_paragraph("维护前必须断电。")
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, [("设备", "周期"), ("泵", "每月"), ("阀门", "每季度")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_paragraph("检查后重新启动。")
    document.add_heading("例行检查", level=2)
    document.add_paragraph()
    document.add_paragraph("记录检查结果。")
    document.add_heading("注意事项", level=3)
    document.add_paragraph("不得带电维护。")
    document.add_heading("停机检查", level=2)
    document.add_paragraph("确认阀门关闭。")
    document.add_heading("设备维护", level=1)
    document.add_paragraph("下轮维护沿用周期表。")
    document.save(SAMPLES / "docx_sections.docx")

    empty = Document()
    empty.add_paragraph(" \t\n ")
    empty.add_heading(" ", level=1)
    empty.add_table(rows=2, cols=2)
    empty.save(SAMPLES / "docx_empty.docx")

    cells = Document()
    cells.add_heading("周期表", level=1)
    table = cells.add_table(rows=3, cols=2)
    for row, values in zip(
        table.rows,
        [("设备 | 部件", "周期"), ("泵\\备用", "每月\n断电检查"), ("阀门", "每季度")],
    ):
        for cell, value in zip(row.cells, values):
            cell.text = value
    table.cell(2, 1).add_paragraph("记录结果")
    cells.save(SAMPLES / "docx_table_cells.docx")

    headings = Document()
    heading_style = headings.styles["Heading 1"]
    heading_style.name = "设备一级标题"
    headings.add_paragraph("维护规程", style=heading_style)
    inherited = headings.styles.add_style("检查标题", WD_STYLE_TYPE.PARAGRAPH)
    inherited.base_style = headings.styles["Heading 2"]
    headings.add_paragraph("例行检查", style=inherited)
    paragraph = headings.add_paragraph("停机规则")
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "2")
    paragraph._p.get_or_add_pPr().append(outline)
    plain = headings.styles.add_style("Heading Notes", WD_STYLE_TYPE.PARAGRAPH)
    headings.add_paragraph("名称含 Heading 的普通说明。", style=plain)
    headings.save(SAMPLES / "docx_outline_headings.docx")

    deep = Document()
    deep.add_heading("六级标题", level=6)
    deep.add_heading("七级标题", level=7)
    deep.add_heading("九级标题", level=9)
    deep.save(SAMPLES / "docx_deep_headings.docx")

    (SAMPLES / "docx_corrupt.docx").write_bytes(b"not a DOCX package")
    with ZipFile(SAMPLES / "docx_teaching.docx") as original:
        for filename, replacement in (
            ("docx_invalid_xml.docx", b"<w:document"),
            ("docx_missing_part.docx", None),
        ):
            with ZipFile(SAMPLES / filename, "w") as damaged:
                for item in original.infolist():
                    if item.filename == "word/document.xml":
                        if replacement is not None:
                            damaged.writestr(copy(item), replacement)
                    else:
                        damaged.writestr(copy(item), original.read(item.filename))


if __name__ == "__main__":
    main()
