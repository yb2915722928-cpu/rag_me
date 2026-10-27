from io import BytesIO
from pathlib import Path
from struct import unpack_from
from zipfile import ZIP_DEFLATED, ZipFile

from PIL import Image
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches


ROOT = Path(__file__).parent


def textbox(shapes, left, top, lines):
    shape = shapes.add_textbox(Inches(left), Inches(top), Inches(3), Inches(1.5))
    frame = shape.text_frame
    for index, (text, level) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.level = level
    return shape


def build_refund():
    presentation = Presentation()
    for number in range(1, 8):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"前置示例 {number}"
        slide.placeholders[1].text = f"本页条件仅属于原第 {number} 页。"
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "退款审批流程"
    textbox(slide.shapes, 6, 1.5, [
        ("企业订单", 0), ("提交申请", 1), ("财务审核", 1), ("退款", 1),
    ])
    textbox(slide.shapes, 0.5, 1.5, [
        ("普通订单", 0), ("提交申请", 1), ("客服审核", 1), ("退款", 1),
    ])
    textbox(slide.shapes, 0.5, 4.5, [("企业订单超过1万元，需负责人追加审批。", 0)])
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(6.5), Inches(3), Inches(6.5), Inches(4),
    )
    arrow = OxmlElement("a:tailEnd")
    arrow.set("type", "triangle")
    connector._element.spPr.get_or_add_ln().append(arrow)
    slide.notes_slide.notes_text_frame.text = "讲解备注：追加审批条件仅适用于企业订单。"
    presentation.save(ROOT / "pptx_refund.pptx")
    (ROOT / "pptx_override.bin").write_bytes((ROOT / "pptx_refund.pptx").read_bytes())


def build_visual():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    textbox(group.shapes, 0.5, 0.5, [("组合内说明", 0), ("#不是标题 <b>原文</b>", 1)])
    group.left = Inches(1.5)
    group.top = Inches(0.75)
    group.width = Inches(4.5)
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4), Inches(1), Inches(2), Inches(1))
    shape.text = "视觉步骤"
    table = slide.shapes.add_table(3, 2, Inches(0.5), Inches(3), Inches(5), Inches(2)).table
    table.cell(0, 0).text = "订单类型"
    table.cell(0, 1).text = "审核部门"
    table.cell(1, 0).text = "企业订单"
    table.cell(1, 1).text = "财务审核"
    table.cell(2, 0).text = "企业专属追加审批"
    table.cell(2, 0).merge(table.cell(2, 1))
    chart = ChartData()
    chart.categories = ["普通订单", "企业订单"]
    chart.add_series("审批步骤", [2, 3])
    chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6), Inches(3), Inches(3), Inches(2), chart)
    workbook = BytesIO(chart_shape.chart.part.chart_workbook.xlsx_part.blob)
    slide.shapes.add_ole_object(workbook, "Excel.Sheet.12", Inches(6), Inches(5.5))
    image = Image.new("RGB", (32, 24), "navy")
    image.save(ROOT / "pptx_visual.png")
    slide.shapes.add_picture(str(ROOT / "pptx_visual.png"), Inches(7), Inches(0.5), Inches(1), Inches(1))
    presentation.save(ROOT / "pptx_visual.pptx")
    image_only = Presentation()
    page = image_only.slides.add_slide(image_only.slide_layouts[6])
    page.shapes.add_picture(str(ROOT / "pptx_visual.png"), Inches(1), Inches(1))
    image_only.save(ROOT / "pptx_image_only.pptx")


def build_tables():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    ordinary = slide.shapes.add_table(2, 2, Inches(0.5), Inches(0.5), Inches(5), Inches(2)).table
    ordinary.cell(0, 0).text = "订单类型"
    ordinary.cell(0, 1).text = "审核部门"
    ordinary.cell(1, 0).text = "普通订单"
    ordinary.cell(1, 1).text = "客服 <审核> & 退款\n原文说明"
    merged = slide.shapes.add_table(2, 2, Inches(0.5), Inches(3), Inches(5), Inches(2)).table
    merged.cell(0, 0).text = "企业订单"
    merged.cell(0, 0).merge(merged.cell(1, 0))
    merged.cell(0, 1).text = "财务审核"
    merged.cell(1, 1).text = "超过1万元追加负责人审批"
    presentation.save(ROOT / "pptx_tables.pptx")


def build_empty_shapes():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    empty = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    empty.name = "空白矩形"
    whitespace = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4), Inches(1), Inches(2), Inches(1))
    whitespace.name = "仅空白字符的矩形"
    whitespace.text = " \t\n\v "
    textbox(slide.shapes, 1, 3, [("", 0), (" \t ", 1)]).name = "空白文本框"
    presentation.save(ROOT / "pptx_empty_shapes.pptx")


def build_entity_literals():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    text = "AT&amp;T &lt;10000 &copy;"
    slide.shapes.title.text = text
    textbox(slide.shapes, 1, 2, [(text, 0)])
    slide.notes_slide.notes_text_frame.text = text
    presentation.save(ROOT / "pptx_entity_literals.pptx")


def build_failures():
    """保留有效前七页；每个输出新开原包，避免 writestr 修改 ZipInfo 后污染其他案例。"""
    source = (ROOT / "pptx_refund.pptx").read_bytes()
    with ZipFile(BytesIO(source)) as archive:
        slide = archive.read("ppt/slides/slide8.xml")
    cases = (
        ("pptx_atomic_failure.pptx", b"<broken"),
        ("pptx_atomic_shape_failure.pptx", slide.replace(b'id="3"', b'id="invalid"', 1)),
        ("pptx_atomic_root_failure.pptx", slide.replace(b"<p:sld ", b"<p:wrong ", 1).replace(b"</p:sld>", b"</p:wrong>", 1)),
    )
    for name, replacement in cases:
        with ZipFile(BytesIO(source)) as archive, ZipFile(ROOT / name, "w", ZIP_DEFLATED) as output:
            for member in archive.infolist():
                data = replacement if member.filename == "ppt/slides/slide8.xml" else archive.read(member.filename)
                output.writestr(member, data)
    with ZipFile(BytesIO(source)) as archive:
        member = archive.getinfo("ppt/slides/slide8.xml")
        name_length, extra_length = unpack_from("<HH", source, member.header_offset + 26)
        compressed_start = member.header_offset + 30 + name_length + extra_length
    damaged = bytearray(source)
    damaged[compressed_start] = 0x07
    (ROOT / "pptx_corrupt_compression.pptx").write_bytes(damaged)


if __name__ == "__main__":
    build_refund()
    build_visual()
    build_tables()
    build_empty_shapes()
    build_entity_literals()
    Presentation().save(ROOT / "pptx_empty.pptx")
    (ROOT / "pptx_corrupt.pptx").write_bytes(b"not a presentation\n")
    build_failures()
