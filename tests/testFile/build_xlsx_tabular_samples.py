"""使用 openpyxl 3.1.5 构造 #16 固定输入；测试运行时不重建。"""
from pathlib import Path

from openpyxl import Workbook


ROOT = Path(__file__).parent
book = Workbook()
sheet = book.active
sheet.title = "差旅"
sheet.append(["城市", "住宿上限（元／晚）", "餐补（元／天）"])
sheet.append(["北京", 500, 100])
sheet.append(["上海", 450, 100])
sheet.append(["杭州", 400, 80])
book.save(ROOT / "xlsx_travel.xlsx")

book = Workbook()
sheet = book.active
sheet.title = "第一表"
sheet.append([None, None])
sheet.append([None, "城市", "金额"])
sheet.append([None, "北京", 500])
sheet.append([None, None, None])
sheet.append([None, "上海", 450])
sheet["H20"].number_format = "0.00"
book.create_sheet("空表")["C8"].number_format = "0.00"
hidden = book.create_sheet("隐藏")
hidden.sheet_state = "hidden"
hidden.append(["城市", "金额"])
hidden.append(["杭州", 400])
second = book.create_sheet("第二表")
second.append(["城市", "金额"])
second.append(["深圳", 300])
book.save(ROOT / "xlsx_sheets.xlsx")

book = Workbook()
sheet = book.active
sheet.title = "类型"
sheet.append(["空白", "零", "布尔", "编号", "小数", "公式", "特殊|列"])
sheet.append([None, 0, False, "00123", 12.5, "=B2+1", "A|B\\路径\n第二行<br>&"])
sheet["E2"].number_format = "0.00"
book.save(ROOT / "xlsx_plain_values.xlsx")

for filename, records in (
    ("xlsx_empty.xlsx", []),
    ("xlsx_header_only.xlsx", [["城市", "金额"]]),
    ("xlsx_duplicate_header.xlsx", [["城市", "城市"], ["北京", 500]]),
    ("xlsx_blank_header.xlsx", [["城市", None], ["北京", 500]]),
):
    book = Workbook()
    for record in records:
        book.active.append(record)
    book.save(ROOT / filename)

book = Workbook()
book.active.title = "有效"
book.active.append(["城市", "金额"])
book.active.append(["北京", 500])
sheet = book.create_sheet("无数据")
sheet.append(["城市", "金额"])
book.save(ROOT / "xlsx_atomic_failure.xlsx")
(ROOT / "xlsx_corrupt.xlsx").write_bytes(b"not an XLSX zip file")

book = Workbook()
sheet = book.active
sheet.title = "合并"
sheet.append(["城市", "金额"])
sheet.append(["北京", 500])
sheet.append([None, 450])
sheet.merge_cells("A2:A3")
book.save(ROOT / "xlsx_merged_unsupported.xlsx")
