"""构造评审固定输入；预设的 OOXML 缓存不代表办公软件重算。"""

from datetime import date
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula


def save_patched(book, name, sheet_path, patch):
    """只改指定工作表的测试 XML，其他工作簿条目原样保留。"""
    buffer = BytesIO()
    book.save(buffer)
    with ZipFile(buffer) as original, ZipFile(Path(__file__).with_name(name), "w") as output:
        for info in original.infolist():
            data = original.read(info.filename)
            if info.filename == sheet_path:
                root = ET.fromstring(data)
                patch(root)
                data = ET.tostring(root, encoding="utf-8")
            output.writestr(info, data)


def string_cache(root):
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    for coordinate, value in (("B2", ""), ("B3", "x")):
        cell = root.find(f'.//s:c[@r="{coordinate}"]', ns)
        cell.set("t", "str")
        cell.find("s:v", ns).text = value


def corrupt_shared_string(root):
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    cell = root.find('.//s:c[@r="A2"]', ns)
    cell.set("t", "s")
    for child in list(cell):
        cell.remove(child)
    ET.SubElement(cell, f"{{{ns['s']}}}v").text = "999"


def build():
    book = Workbook()
    sheet = book.active
    sheet.append(["类型", "值"])
    sheet.append(["空文本缓存", '=IF(1=1,"","x")'])
    sheet.append(["文本缓存", '="x"'])
    sheet.append(["无缓存", "=1+1"])
    save_patched(book, "xlsx_review_string_cache.xlsx", "xl/worksheets/sheet1.xml", string_cache)

    book = Workbook()
    sheet = book.active
    sheet.append(["对象", "结果"])
    sheet.append(["A", None])
    sheet["B2"] = ArrayFormula(ref="B2:B2", text='=SUM(IF(A2="A",1,0))')
    book.save(Path(__file__).with_name("xlsx_review_array_formula.xlsx"))

    book = Workbook()
    sheet = book.active
    sheet.append(["项目", "金额"])
    sheet.merge_cells("A2:A3")
    sheet["A2"] = "   "
    book.save(Path(__file__).with_name("xlsx_review_blank_merge_no_data.xlsx"))
    sheet["A4"] = "有效"
    sheet["B4"] = 0
    book.save(Path(__file__).with_name("xlsx_review_blank_merge.xlsx"))

    book = Workbook()
    sheet = book.active
    sheet.title = "有效表"
    sheet.append(["对象", "值"])
    sheet.append(["有效内容", 1])
    sheet = book.create_sheet("损坏表")
    sheet.append(["对象", "值"])
    sheet.append(["损坏内容", 2])
    save_patched(book, "xlsx_review_corrupt_shared_string.xlsx",
                 "xl/worksheets/sheet2.xml", corrupt_shared_string)

    book = Workbook(iso_dates=True)
    sheet = book.active
    sheet.append(["项目", "日期"])
    sheet.append(["ISO 日期", date(2026, 10, 27)])
    sheet["B2"].number_format = "yyyy-mm-dd"
    book.save(Path(__file__).with_name("xlsx_review_iso_date.xlsx"))


if __name__ == "__main__":
    build()
