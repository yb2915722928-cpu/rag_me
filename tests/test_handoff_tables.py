import copy
from datetime import date, timedelta
import importlib
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook


tabular = importlib.import_module(
    "app.modules.1_documentIngestionAndParsing.4_tabularFileProcessing"
)
handoff = importlib.import_module(
    "app.modules.1_documentIngestionAndParsing.unifiedDocumentHandoff"
)
SAMPLES = Path(__file__).parent / "testFile"
CONTEXT = tabular.TableContext(
    "员工国内差旅标准", "2026示例版", "普通员工，国内出差", "教学数据，不是实际报销政策",
)


class TableHandoffTests(unittest.TestCase):
    def export(self, report):
        with tempfile.TemporaryDirectory() as directory:
            output = handoff.export_json(report, Path(directory) / "表格.json")
            return json.loads(output.read_bytes())

    def test_export_rejects_unrecognized_objects_instead_of_silently_stringifying_them(self):
        native = tabular.convert_file(SAMPLES / "csv_travel.csv", context=CONTEXT)
        report = handoff.adapt_result(native)
        report.metadata["unsupported"] = object()
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "交接.json"
            output.write_text("existing", encoding="utf-8")
            with self.assertRaises(TypeError):
                handoff.export_json(report, output)
            self.assertEqual(output.read_text(encoding="utf-8"), "existing")

    def test_merged_measurements_do_not_turn_covered_blanks_into_independent_values(self):
        native = tabular.convert_file(SAMPLES / "xlsx_structure_measurements.xlsx", context=CONTEXT)
        self.assertTrue(native.success, native.error)
        before = copy.deepcopy(native)
        document = self.export(handoff.adapt_result(native))["documents"][0]
        anchor, covered, outside = document["structure"]["tables"][0]["rows"]
        self.assertEqual([cell["raw_value"] for cell in anchor["cells"]], ["甲", 900, "=B2+1", True])
        self.assertEqual([cell["raw_value"] for cell in covered["cells"]], ["乙", None, None, None])
        for cell, coordinate in zip(covered["cells"][1:], ("B2", "C2", "D2")):
            self.assertEqual(cell["source"]["merged_anchor"], coordinate)
            self.assertIsNone(cell["formula"])
            self.assertIsNone(cell["cached_value"])
            self.assertIsNone(cell["cache_status"])
            self.assertIn("非本行独立值", cell["text"])
        self.assertEqual([cell["text"] for cell in outside["cells"][1:]], ["", "", ""])
        self.assertIn("非本行独立值", document["units"][1]["text"])
        self.assertNotIn("900", document["units"][1]["text"])
        self.assertEqual(document["body"], native.markdown)
        self.assertEqual(native, before)

    def test_saved_empty_text_cache_stays_distinct_from_an_unknown_formula_result(self):
        native = tabular.convert_file(SAMPLES / "xlsx_review_string_cache.xlsx", context=CONTEXT)
        self.assertTrue(native.success, native.error)
        before = copy.deepcopy(native)
        document = self.export(handoff.adapt_result(native))["documents"][0]
        empty, text, unknown = [row["cells"][1] for row in document["structure"]["tables"][0]["rows"]]
        self.assertEqual(empty["formula"], '=IF(1=1,"","x")')
        self.assertEqual((empty["cached_value"], empty["cache_status"]), ("", "cached"))
        self.assertIn("空文本", empty["text"])
        self.assertEqual((text["cached_value"], text["cache_status"]), ("x", "cached"))
        self.assertEqual((unknown["cached_value"], unknown["cache_status"]), (None, "unknown"))
        self.assertIn("结果未知", unknown["text"])
        self.assertEqual(document["body"], native.markdown)
        self.assertEqual(native, before)

    def test_native_date_objects_export_as_dates_and_projection_is_independent(self):
        native = tabular.convert_file(SAMPLES / "xlsx_review_iso_date.xlsx", context=CONTEXT)
        self.assertTrue(native.success, native.error)
        before = copy.deepcopy(native)
        self.assertIs(type(native.tables[0].rows[0].cells[1].raw_value), date)
        report = handoff.adapt_result(native)
        projected = report.to_dict()
        document = self.export(report)["documents"][0]
        cell = document["structure"]["tables"][0]["rows"][0]["cells"][1]
        self.assertEqual(cell["raw_value"], "2026-10-27")
        self.assertEqual(cell["text"], "2026-10-27")
        self.assertEqual(document["body"], native.markdown)
        self.assertEqual(document["diagnostics"], [])
        projected["documents"][0]["structure"]["tables"][0]["rows"][0]["cells"][1]["raw_value"] = "改变"
        self.assertEqual(report.documents[0].structure["tables"][0]["rows"][0]["cells"][1]["raw_value"],
                         date(2026, 10, 27))
        self.assertEqual(native, before)

    def test_xlsx_multiple_sheets_keep_rows_and_cells_attached_to_the_original_tables(self):
        path = SAMPLES / "xlsx_sheets.xlsx"
        native = tabular.convert_bytes(path.read_bytes(), format="xlsx", source="多表.xlsx", context=CONTEXT)
        self.assertTrue(native.success, native.error)
        before = copy.deepcopy(native)
        document = self.export(handoff.adapt_result(native))["documents"][0]
        tables = document["structure"]["tables"]
        self.assertEqual([table["source"]["sheet"] for table in tables], ["第一表", "第二表"])
        self.assertEqual(tables[0]["source"]["coordinate"], "B2:C5")
        self.assertEqual(tables[0]["columns"][1]["source"]["coordinate"], "C2")
        self.assertEqual(tables[0]["rows"][1]["cells"][1]["raw_value"], 450)
        self.assertEqual(tables[0]["rows"][1]["cells"][1]["source"]["coordinate"], "C5")
        self.assertEqual(tables[1]["rows"][0]["cells"][0]["raw_value"], "深圳")
        units = document["units"]
        self.assertEqual([(unit["metadata"]["table_index"], unit["metadata"]["row_index"]) for unit in units],
                         [(0, 0), (0, 1), (1, 0)])
        sources = {source["id"]: source for source in document["sources"]}
        positions = [sources[unit["source_ids"][0]]["locator"] for unit in units]
        self.assertEqual([(position["sheet"], position["coordinate"]) for position in positions],
                         [("第一表", "B3:C3"), ("第一表", "B5:C5"), ("第二表", "A2:B2")])
        self.assertTrue(all(position["source"]["location"] == "多表.xlsx" for position in positions))
        self.assertTrue(all(unit["body_span"] is None and unit["alignment"] == "unmapped" for unit in units))
        all_ids = [document["id"], *sources, *(unit["id"] for unit in units)]
        self.assertEqual(len(all_ids), len(set(all_ids)))
        self.assertEqual(document["body"], native.markdown)
        self.assertEqual(native, before)

    def test_xlsx_merged_headers_and_covered_cells_keep_units_context_and_real_anchors(self):
        options = tabular.XlsxOptions((tabular.XlsxSheetOptions("合并", header_rows=(1, 2, 3)),))
        for mode in ("table", "rows"):
            with self.subTest(mode=mode):
                native = tabular.convert_file(
                    SAMPLES / "xlsx_structure_merges.xlsx", xlsx=options, context=CONTEXT, mode=mode,
                )
                self.assertTrue(native.success, native.error)
                before = copy.deepcopy(native)
                document = self.export(handoff.adapt_result(native))["documents"][0]
                self.assertEqual(document["body"], native.markdown)
                table = document["structure"]["tables"][0]
                self.assertEqual([column["name"] for column in table["columns"]],
                                 ["地区", "差旅标准 / 住宿 / 元／晚", "差旅标准 / 餐补 / 元／天", "备注"])
                header_sources = table["columns"][2]["header_sources"]
                self.assertEqual([source["coordinate"] for source in header_sources], ["C1", "C2", "C3"])
                self.assertEqual((header_sources[0]["merged_anchor"], header_sources[0]["merged_range"]),
                                 ("B1", "B1:C1"))
                self.assertEqual(table["context"]["notes"], "教学数据，不是实际报销政策")
                first, covered, outside, horizontal = table["rows"]
                self.assertEqual(first["cells"][0]["raw_value"], "华北")
                self.assertIsNone(covered["cells"][0]["raw_value"])
                self.assertEqual(covered["cells"][0]["text"], "华北（合并共同标签，锚点 A4，范围 A4:A5）")
                self.assertEqual(covered["cells"][0]["source"]["coordinate"], "A5")
                self.assertEqual(covered["cells"][0]["source"]["merged_anchor"], "A4")
                self.assertEqual(covered["cells"][3]["source"]["merged_range"], "D4:D5")
                self.assertIsNone(outside["cells"][0]["raw_value"])
                self.assertEqual(outside["cells"][0]["text"], "")
                self.assertIsNone(outside["cells"][0]["source"]["merged_anchor"])
                self.assertIsNone(horizontal["cells"][2]["raw_value"])
                self.assertEqual(horizontal["cells"][2]["source"]["merged_range"], "B7:D7")
                sources = {source["id"]: source for source in document["sources"]}
                self.assertEqual([unit["text"] for unit in document["units"]],
                                 [row.text for row in native.tables[0].rows])
                self.assertIn("差旅标准 / 住宿 / 元／晚：500", document["units"][0]["text"])
                for unit in document["units"]:
                    self.assertEqual(unit["alignment"], "unmapped")
                    self.assertIsNone(unit["body_span"])
                    source = sources[unit["source_ids"][0]]["locator"]
                    self.assertEqual(source["sheet"], "合并")
                    self.assertIsNone(source["line_start"])
                    self.assertIsNone(source["line_end"])
                self.assertEqual(native, before)

    def test_xlsx_durations_export_signed_iso_seconds_instead_of_clock_times(self):
        book = load_workbook(SAMPLES / "xlsx_values.xlsx")
        for coordinate, duration in (
            ("B16", timedelta(days=1, hours=2, minutes=3, seconds=4, milliseconds=5)),
            ("B17", -timedelta(hours=1, milliseconds=5)),
            ("B18", timedelta(0)),
        ):
            book.active[coordinate] = duration
            book.active[coordinate].number_format = "[h]:mm:ss"
        buffer = BytesIO()
        book.save(buffer)
        book.close()
        native = tabular.convert_bytes(buffer.getvalue(), format="xlsx", source="持续时间.xlsx", context=CONTEXT)
        self.assertTrue(native.success, native.error)
        before = copy.deepcopy(native)
        self.assertIsInstance(native.tables[0].rows[14].cells[1].raw_value, timedelta)

        report = handoff.adapt_result(native)
        document = self.export(report)["documents"][0]
        cells = {row["cells"][1]["source"]["coordinate"]: row["cells"][1]
                 for row in document["structure"]["tables"][0]["rows"]}
        self.assertEqual(cells["B16"]["raw_value"], "PT93784.005S")
        self.assertEqual(cells["B17"]["raw_value"], "-PT3600.005S")
        self.assertEqual(cells["B18"]["raw_value"], "PT0S")
        self.assertEqual(cells["B16"]["number_format"], "[h]:mm:ss")
        self.assertEqual(document["body"], native.markdown)
        self.assertEqual([unit["text"] for unit in document["units"]],
                         [row.text for row in native.tables[0].rows])
        self.assertEqual(native, before)

    def test_xlsx_values_export_iso_dates_without_changing_scalars_formula_caches_or_input(self):
        path = SAMPLES / "xlsx_values.xlsx"
        for mode in ("table", "rows"):
            for entry in ("file", "bytes"):
                with self.subTest(mode=mode, entry=entry):
                    if entry == "file":
                        native = tabular.convert_file(path, context=CONTEXT, mode=mode)
                    else:
                        native = tabular.convert_bytes(
                            path.read_bytes(), format="xlsx", source=str(path), context=CONTEXT, mode=mode,
                        )
                    self.assertTrue(native.success, native.error)
                    before = copy.deepcopy(native)
                    report = handoff.adapt_result(native)
                    report_before = copy.deepcopy(report)
                    document = self.export(report)["documents"][0]

                    self.assertEqual(document["body"], native.markdown)
                    self.assertEqual([unit["text"] for unit in document["units"]],
                                     [row.text for row in native.tables[0].rows])
                    cells = {row["cells"][1]["source"]["coordinate"]: row["cells"][1]
                             for row in document["structure"]["tables"][0]["rows"]}
                    for coordinate, raw, text in (
                        ("B16", "2026-10-27T00:00:00", "2026-10-27"),
                        ("B17", "2026-10-27T13:04:05", "2026-10-27T13:04:05"),
                        ("B18", "13:04:05", "13:04:05"),
                    ):
                        self.assertEqual(cells[coordinate]["raw_value"], raw)
                        self.assertEqual(cells[coordinate]["text"], text)
                        self.assertEqual(cells[coordinate]["data_type"], "d")
                    self.assertEqual(cells["B16"]["number_format"], "yyyy-mm-dd")
                    for coordinate, raw in (("B6", None), ("B7", 0), ("B12", 0.2),
                                            ("B22", False), ("B23", "00123")):
                        self.assertEqual(cells[coordinate]["raw_value"], raw)
                        self.assertIs(type(cells[coordinate]["raw_value"]), type(raw))
                    for coordinate, formula, cached, status in (
                        ("B2", "=1-1", 0, "cached"),
                        ("B3", "=1+1", None, "unknown"),
                        ("B4", "=40+2", 42, "cached"),
                        ("B5", "=1/0", "#DIV/0!", "cached"),
                        ("B20", "=DATE(2026,10,27)", "2026-10-27T00:00:00", "cached"),
                        ("B26", '=IF(1,"00123","")', "00123", "cached"),
                        ("B27", "=1=2", False, "cached"),
                    ):
                        cell = cells[coordinate]
                        self.assertEqual(cell["raw_value"], formula)
                        self.assertEqual(cell["formula"], formula)
                        self.assertEqual(cell["cached_value"], cached)
                        self.assertIs(type(cell["cached_value"]), type(cached))
                        self.assertEqual(cell["cache_status"], status)
                    self.assertIn("非实时计算结果", cells["B2"]["text"])
                    self.assertIn("结果未知", cells["B3"]["text"])
                    self.assertEqual([item["message"] for item in document["diagnostics"]], native.warnings)
                    self.assertEqual(report.to_dict()["documents"][0]["structure"]["tables"][0]["rows"],
                                     document["structure"]["tables"][0]["rows"])
                    self.assertEqual(native, before)
                    self.assertEqual(report, report_before)

    def test_upstream_failures_and_unknown_context_keep_diagnostics_outside_body(self):
        for filename, success, options in (
            ("csv_travel.csv", True, {}),
            ("csv_explicit_gb18030.csv", True, {"delimiter": ";", "header_record": 3}),
            ("csv_quote_unclosed.csv", False, {}),
            ("xlsx_corrupt.xlsx", False, {}),
        ):
            with self.subTest(filename=filename):
                native = tabular.convert_file(SAMPLES / filename, **options)
                self.assertEqual(native.success, success, native.error)
                before = copy.deepcopy(native)
                report = handoff.adapt_result(native)
                with tempfile.TemporaryDirectory() as directory:
                    output = Path(directory) / "诊断.json"
                    output.write_text("existing", encoding="utf-8")
                    saved = json.loads(handoff.export_json(report, output).read_bytes())
                document = saved["documents"][0]
                self.assertEqual(saved["status"], "success" if success else "failed")
                self.assertEqual(document["status"], saved["status"])
                diagnostics = document["diagnostics"]
                self.assertEqual([item["message"] for item in diagnostics if item["level"] == "warning"],
                                 native.warnings)
                source = document["sources"][0]
                self.assertEqual(source["locator"]["location"], str(SAMPLES / filename))
                self.assertEqual(document["metadata"]["encoding"], native.encoding)
                for diagnostic in diagnostics:
                    self.assertEqual(diagnostic["document_id"], document["id"])
                    self.assertEqual(diagnostic["source_id"], source["id"])
                    self.assertEqual(diagnostic["code"], "upstream_" + diagnostic["level"])
                if success:
                    self.assertEqual(document["body"], native.markdown)
                    context = document["structure"]["tables"][0]["context"]
                    self.assertEqual(context, {"name": None, "version": None, "applicability": None, "notes": None})
                    self.assertIn("表名：未知\n版本：未知\n适用条件：未知", document["units"][0]["text"])
                    for warning in native.warnings:
                        self.assertNotIn(warning, document["body"])
                        self.assertNotIn(warning, document["units"][0]["text"])
                else:
                    self.assertIsNone(document["body"])
                    self.assertEqual(document["units"], [])
                    self.assertEqual(document["structure"]["tables"], [])
                    self.assertEqual(diagnostics[-1]["message"], native.error)
                    self.assertEqual(diagnostics[-1]["level"], "error")
                self.assertIsNone(document["business_id"])
                self.assertIsNone(document["version"])
                self.assertFalse(document["requires_review"])
                self.assertEqual(native, before)

    def test_csv_structure_links_rows_to_columns_cells_and_original_record_positions(self):
        native = tabular.convert_file(SAMPLES / "csv_quoted.csv", mode="table", context=CONTEXT)
        self.assertTrue(native.success, native.error)
        before = copy.deepcopy(native)
        report = handoff.adapt_result(native, business_id="差旅标准", version="修订二")
        document = self.export(report)["documents"][0]

        self.assertEqual(document["business_id"], "差旅标准")
        self.assertEqual(document["version"], "修订二")
        tables = document["structure"]["tables"]
        self.assertEqual(len(tables), 1)
        table = tables[0]
        self.assertEqual(table["context"]["applicability"], "普通员工，国内出差")
        self.assertEqual(table["context"]["notes"], "教学数据，不是实际报销政策")
        self.assertEqual([column["name"] for column in table["columns"]], ["编号", "说明", "金额", "备注"])
        self.assertEqual(table["columns"][1]["source"]["column_start"], 2)
        self.assertEqual(table["columns"][1]["source"]["row_start"], 2)
        first, last = table["rows"]
        self.assertEqual([cell["raw_value"] for cell in first["cells"]],
                         ["00123", '含逗号,与"引号"\r\n第二行', "=1+2", ""])
        self.assertIsNone(first["cells"][2]["formula"])
        self.assertEqual(first["cells"][1]["source"]["line_end"], 4)
        self.assertEqual(first["columns"], table["columns"])
        self.assertEqual(first["context"], table["context"])
        self.assertEqual([cell["raw_value"] for cell in last["cells"]], ["00002", "普通文本", "0", "尾"])
        sources = {source["id"]: source for source in document["sources"]}
        for row_index, unit in enumerate(document["units"]):
            self.assertEqual(unit["metadata"], {"table_index": 0, "row_index": row_index})
            row = tables[unit["metadata"]["table_index"]]["rows"][unit["metadata"]["row_index"]]
            self.assertEqual(unit["text"], row["text"])
            self.assertEqual(len(unit["source_ids"]), 1)
            source = sources[unit["source_ids"][0]]
            self.assertEqual(source["kind"], "table_row")
            self.assertEqual(source["locator"], row["source"])
            self.assertIsNone(source["locator"]["sheet"])
            self.assertIsNone(source["locator"]["coordinate"])
        first_source = sources[document["units"][0]["source_ids"][0]]["locator"]
        last_source = sources[document["units"][1]["source_ids"][0]]["locator"]
        self.assertEqual((first_source["row_start"], first_source["row_end"],
                          first_source["line_start"], first_source["line_end"]), (3, 3, 3, 4))
        self.assertEqual((last_source["row_start"], last_source["line_start"]), (6, 7))
        all_ids = [document["id"], *sources, *(unit["id"] for unit in document["units"])]
        self.assertEqual(len(all_ids), len(set(all_ids)))
        self.assertEqual(native, before)
        report.documents[0].structure["tables"][0]["rows"][0]["cells"][0]["raw_value"] = "改变"
        report.documents[0].sources[1].locator["source"]["location"] = "改变"
        self.assertEqual(native, before)

    def test_csv_handoff_keeps_original_body_and_upstream_rows_in_both_modes(self):
        path = SAMPLES / "csv_travel.csv"
        for mode in ("table", "rows"):
            for entry in ("file", "bytes"):
                with self.subTest(mode=mode, entry=entry):
                    if entry == "file":
                        native = tabular.convert_file(path, context=CONTEXT, mode=mode)
                    else:
                        native = tabular.convert_bytes(
                            path.read_bytes(), format="csv", source=str(path), context=CONTEXT, mode=mode,
                        )
                    self.assertTrue(native.success, native.error)
                    before = copy.deepcopy(native)
                    saved = self.export(handoff.adapt_result(native))

                    self.assertEqual(saved["status"], "success")
                    self.assertEqual(len(saved["documents"]), 1)
                    document = saved["documents"][0]
                    self.assertEqual(document["body"], native.markdown)
                    self.assertEqual(document["body_format"], "markdown")
                    self.assertEqual(document["metadata"]["encoding"], native.encoding)
                    self.assertIsNone(document["business_id"])
                    self.assertIsNone(document["version"])
                    units = document["units"]
                    self.assertEqual([unit["text"] for unit in units],
                                     [row.text for row in native.tables[0].rows])
                    self.assertEqual(len(units), 3)
                    self.assertIn("城市：北京\n住宿上限（元／晚）：500", units[0]["text"])
                    self.assertIn("城市：上海\n住宿上限（元／晚）：450", units[1]["text"])
                    for unit in units:
                        self.assertEqual(unit["kind"], "table_row")
                        self.assertEqual(unit["alignment"], "unmapped")
                        self.assertIsNone(unit["body_span"])
                        for evidence in ("员工国内差旅标准", "2026示例版", "普通员工，国内出差", "教学数据"):
                            self.assertIn(evidence, unit["text"])
                    self.assertEqual(native, before)


if __name__ == "__main__":
    unittest.main()
